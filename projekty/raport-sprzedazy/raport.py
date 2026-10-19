"""
Automatyczny raport sprzedaży
=============================

Wrzuć pliki z oddziałów (.xlsx / .xls / .csv) do folderu `dane_wejsciowe/`
i uruchom:

    python raport.py

Skrypt:
  1. wczytuje wszystkie pliki, niezależnie od nazw kolumn i formatu,
  2. ujednolica daty, kwoty i nazwy produktów, usuwa duplikaty i puste wiersze,
  3. zapisuje gotowy raport `raport_sprzedazy.xlsx` z podsumowaniem,
     wykresami, rankingiem handlowców, czystymi danymi i logiem poprawek.

Wszystkie liczby w raporcie to formuły Excela liczone z arkusza "Dane",
więc raport da się dalej edytować i filtrować.
"""
from __future__ import annotations

import difflib
import re
import sys
import unicodedata
from collections import Counter
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.table import Table, TableStyleInfo

# ---------------------------------------------------------------------------
# KONFIGURACJA – to jedyne miejsce, które zmieniasz pod nowego klienta
# ---------------------------------------------------------------------------
FOLDER_WEJSCIE = Path(__file__).parent / "dane_wejsciowe"
PLIK_WYJSCIE = Path(__file__).parent / "raport_sprzedazy.xlsx"
NAZWA_FIRMY = "Kawiarnia & Palarnia „Ziarno”"

# Jak rozpoznać oddział po nazwie pliku (fragment nazwy -> oddział)
ODDZIALY = {"poznan": "Poznań", "waw": "Warszawa", "warszawa": "Warszawa", "wroclaw": "Wrocław"}

# Oficjalny katalog produktów – literówki zostaną dopasowane do tych nazw
KATALOG = [
    "Kawa ziarnista 1kg",
    "Kawa mielona 250g",
    "Herbata zielona",
    "Ekspres przelewowy",
    "Młynek ręczny",
    "Kubek termiczny",
]

# Synonimy nazw kolumn spotykane w plikach oddziałów
KOLUMNY = {
    "data": ["data", "data sprzedaży", "data sprzedazy", "dt", "date"],
    "produkt": ["produkt", "nazwa towaru", "towar", "artykuł", "artykul", "product"],
    "ilosc": ["ilość", "ilosc", "szt", "sztuk", "qty", "quantity"],
    "cena": ["cena", "cena jedn.", "cena jednostkowa", "price"],
    "wartosc": ["wartość", "wartosc", "kwota", "suma", "value", "amount"],
    "handlowiec": ["handlowiec", "sprzedawca", "osoba", "pracownik", "seller"],
}

# ---------------------------------------------------------------------------
# CZYSZCZENIE DANYCH
# ---------------------------------------------------------------------------
log: list[tuple[str, str, int]] = []  # (plik, co poprawiono, ile wierszy)


def bez_ogonkow(tekst: str) -> str:
    tekst = tekst.replace("ł", "l").replace("Ł", "L")
    return "".join(c for c in unicodedata.normalize("NFKD", tekst) if not unicodedata.combining(c))


def klucz(tekst: str) -> str:
    """Uproszczona postać tekstu do porównań: małe litery, bez ogonków, bez spacji i kropek."""
    return re.sub(r"[\s.]", "", bez_ogonkow(str(tekst)).lower())


def wczytaj(plik: Path) -> pd.DataFrame:
    if plik.suffix.lower() == ".csv":
        # automatyczne wykrycie separatora (; albo ,)
        return pd.read_csv(plik, sep=None, engine="python", encoding="utf-8-sig", dtype=str)
    return pd.read_excel(plik, dtype=str)


def mapuj_kolumny(df: pd.DataFrame, plik: str) -> pd.DataFrame:
    nowe = {}
    for kol in df.columns:
        k = klucz(kol)
        for docelowa, synonimy in KOLUMNY.items():
            if k in {klucz(s) for s in synonimy}:
                nowe[kol] = docelowa
    df = df.rename(columns=nowe)
    brak = {"data", "produkt", "ilosc"} - set(df.columns)
    if brak:
        raise ValueError(f"{plik}: nie rozpoznano kolumn {brak}. Dodaj synonim w KOLUMNY.")
    return df[[c for c in KOLUMNY if c in df.columns]]


def parsuj_date(wartosc):
    if pd.isna(wartosc):
        return pd.NaT
    s = str(wartosc).strip().split(" ")[0]
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    return pd.NaT


def parsuj_liczbe(wartosc):
    """'1 234,50 zł' -> 1234.5 ; '18,5' -> 18.5 ; '3.0' -> 3.0"""
    if pd.isna(wartosc):
        return float("nan")
    s = re.sub(r"[^\d,.\-]", "", str(wartosc))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return float("nan")


_KATALOG_KLUCZE = {klucz(p): p for p in KATALOG}


def dopasuj_produkt(nazwa):
    if pd.isna(nazwa):
        return None
    k = klucz(nazwa)
    if k in _KATALOG_KLUCZE:
        return _KATALOG_KLUCZE[k]
    podobne = difflib.get_close_matches(k, _KATALOG_KLUCZE, n=1, cutoff=0.75)
    return _KATALOG_KLUCZE[podobne[0]] if podobne else None


def oczysc(plik: Path) -> pd.DataFrame:
    nazwa = plik.name
    df = mapuj_kolumny(wczytaj(plik), nazwa)

    przed = len(df)
    df = df.dropna(how="all")
    if przed - len(df):
        log.append((nazwa, "Usunięto puste wiersze", przed - len(df)))

    # oddział z nazwy pliku
    oddzial = next((o for frag, o in ODDZIALY.items() if frag in klucz(nazwa)), plik.stem)
    df["oddzial"] = oddzial

    # daty
    df["data"] = df["data"].map(parsuj_date)
    zle = df["data"].isna().sum()
    if zle:
        log.append((nazwa, "Odrzucono wiersze z nieczytelną datą", int(zle)))
    df = df.dropna(subset=["data"])

    # liczby
    df["ilosc"] = df["ilosc"].map(parsuj_liczbe)
    for kol in ("cena", "wartosc"):
        if kol in df.columns:
            df[kol] = df[kol].map(parsuj_liczbe)
    zle = (df["ilosc"].isna() | (df["ilosc"] <= 0)).sum()
    if zle:
        log.append((nazwa, "Odrzucono wiersze z pustą lub zerową ilością", int(zle)))
    df = df[df["ilosc"] > 0]

    # brakująca wartość albo cena – dopełniamy z drugiej kolumny
    if "wartosc" not in df.columns:
        df["wartosc"] = df["ilosc"] * df["cena"]
        log.append((nazwa, "Wyliczono brakującą kolumnę wartości (ilość × cena)", len(df)))
    if "cena" not in df.columns:
        df["cena"] = df["wartosc"] / df["ilosc"]
        log.append((nazwa, "Wyliczono brakującą cenę jednostkową (wartość ÷ ilość)", len(df)))

    # produkty
    oryginal = df["produkt"].copy()
    df["produkt"] = df["produkt"].map(dopasuj_produkt)
    poprawione = (oryginal.str.strip() != df["produkt"]) & df["produkt"].notna()
    if poprawione.sum():
        log.append((nazwa, "Ujednolicono nazwy produktów (literówki, wielkość liter)", int(poprawione.sum())))
    nieznane = df["produkt"].isna().sum()
    if nieznane:
        log.append((nazwa, "Odrzucono produkty spoza katalogu", int(nieznane)))
        print(f"  ! {nazwa}: nieznane produkty: {sorted(set(oryginal[df['produkt'].isna()]))}")
    df = df.dropna(subset=["produkt"])

    df["handlowiec"] = df.get("handlowiec", pd.Series("brak", index=df.index)).fillna("brak").str.strip()

    przed = len(df)
    df = df.drop_duplicates()
    if przed - len(df):
        log.append((nazwa, "Usunięto zduplikowane wiersze", przed - len(df)))

    log.append((nazwa, "Wierszy po czyszczeniu", len(df)))
    return df


# ---------------------------------------------------------------------------
# RAPORT EXCEL
# ---------------------------------------------------------------------------
FONT = "Arial"
GRANAT = "1F3A5F"
JASNY = "EAF0F6"
f_tytul = Font(name=FONT, size=18, bold=True, color=GRANAT)
f_pod = Font(name=FONT, size=10, italic=True, color="666666")
f_nagl = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_sekcja = Font(name=FONT, size=12, bold=True, color=GRANAT)
f_norm = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_kpi = Font(name=FONT, size=16, bold=True, color=GRANAT)
f_kpi_et = Font(name=FONT, size=9, color="666666")
wyp_nagl = PatternFill("solid", fgColor=GRANAT)
wyp_jasny = PatternFill("solid", fgColor=JASNY)
ramka = Border(bottom=Side(style="thin", color="C9D3DE"))
ZL = '#,##0.00 "zł"'
ZL0 = '#,##0 "zł"'
SZT = '#,##0'
PROC = '0.0%'


def naglowek(ws, wiersz, kolumna, teksty):
    for i, t in enumerate(teksty):
        c = ws.cell(wiersz, kolumna + i, t)
        c.font, c.fill = f_nagl, wyp_nagl
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def styl_wiersza(ws, wiersz, k1, k2, pogrub=False):
    for k in range(k1, k2 + 1):
        c = ws.cell(wiersz, k)
        c.font = f_bold if pogrub else f_norm
        c.border = ramka


def zbuduj_raport(dane: pd.DataFrame):
    wb = Workbook()
    ws = wb.active
    ws.title = "Podsumowanie"
    wd = wb.create_sheet("Dane")
    wh = wb.create_sheet("Handlowcy")
    wl = wb.create_sheet("Log czyszczenia")

    # ---------------- Arkusz: Dane ----------------
    dane = dane.sort_values(["data", "oddzial"]).reset_index(drop=True)
    kol_dane = ["Data", "Miesiąc", "Oddział", "Produkt", "Ilość", "Cena jedn.", "Wartość", "Handlowiec"]
    naglowek(wd, 1, 1, kol_dane)
    for i, r in enumerate(dane.itertuples(), start=2):
        wd.cell(i, 1, r.data.date()).number_format = "yyyy-mm-dd"
        wd.cell(i, 2, r.data.strftime("%Y-%m"))
        wd.cell(i, 3, r.oddzial)
        wd.cell(i, 4, r.produkt)
        wd.cell(i, 5, r.ilosc).number_format = SZT
        wd.cell(i, 6, round(r.cena, 2)).number_format = ZL
        wd.cell(i, 7, f"=E{i}*F{i}").number_format = ZL
        wd.cell(i, 8, r.handlowiec)
        for k in range(1, 9):
            wd.cell(i, k).font = f_norm
    ostatni = len(dane) + 1
    tab = Table(displayName="Sprzedaz", ref=f"A1:H{ostatni}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    wd.add_table(tab)
    for k, szer in zip("ABCDEFGH", [12, 10, 12, 22, 8, 12, 14, 22]):
        wd.column_dimensions[k].width = szer
    wd.freeze_panes = "A2"

    # zakresy do formuł
    R = lambda kol: f"Dane!${kol}$2:${kol}${ostatni}"  # noqa: E731
    WART, ILOSC, ODDZ, PROD, MIES, HAND = R("G"), R("E"), R("C"), R("D"), R("B"), R("H")

    # ---------------- Arkusz: Podsumowanie ----------------
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    for k, szer in zip("BCDEFGH", [24, 16, 16, 16, 16, 16, 4]):
        ws.column_dimensions[k].width = szer

    ws["B2"] = f"Raport sprzedaży – {NAZWA_FIRMY}"
    ws["B2"].font = f_tytul
    od, do = dane["data"].min(), dane["data"].max()
    ws["B3"] = (f"Okres: {od:%d.%m.%Y} – {do:%d.%m.%Y}  ·  Wygenerowano automatycznie "
                f"{datetime.now():%d.%m.%Y %H:%M}  ·  Źródło: {dane['oddzial'].nunique()} oddziały")
    ws["B3"].font = f_pod

    # KPI
    kpi = [
        ("Przychód łącznie", f"=SUM({WART})", ZL0),
        ("Liczba transakcji", f"=COUNTA({PROD})", SZT),
        ("Sprzedane sztuki", f"=SUM({ILOSC})", SZT),
        ("Średni paragon", "=B6/C6", ZL),
    ]
    for i, (et, formula, fmt) in enumerate(kpi):
        kol = 2 + i
        e = ws.cell(5, kol, et)
        e.font, e.fill = f_kpi_et, wyp_jasny
        v = ws.cell(6, kol, formula)
        v.font, v.fill, v.number_format = f_kpi, wyp_jasny, fmt
        for c in (e, v):
            c.alignment = Alignment(horizontal="left", indent=1)
    ws.row_dimensions[6].height = 28

    # Oddziały
    oddzialy = sorted(dane["oddzial"].unique())
    w = 9
    ws.cell(w, 2, "Sprzedaż wg oddziałów").font = f_sekcja
    naglowek(ws, w + 1, 2, ["Oddział", "Przychód", "Transakcje", "Sztuki", "Średni paragon", "Udział"])
    pierwszy_odd = w + 2
    for j, o in enumerate(oddzialy):
        r = pierwszy_odd + j
        ws.cell(r, 2, o)
        ws.cell(r, 3, f'=SUMIFS({WART},{ODDZ},B{r})').number_format = ZL0
        ws.cell(r, 4, f'=COUNTIFS({ODDZ},B{r})').number_format = SZT
        ws.cell(r, 5, f'=SUMIFS({ILOSC},{ODDZ},B{r})').number_format = SZT
        ws.cell(r, 6, f'=IFERROR(C{r}/D{r},0)').number_format = ZL
        ws.cell(r, 7, f'=IFERROR(C{r}/$B$6,0)').number_format = PROC
        styl_wiersza(ws, r, 2, 7)
    ostatni_odd = pierwszy_odd + len(oddzialy) - 1
    r = ostatni_odd + 1
    ws.cell(r, 2, "Razem")
    for k, fmt in zip("CDEFG", [ZL0, SZT, SZT, ZL, PROC]):
        ws[f"{k}{r}"] = f"=SUM({k}{pierwszy_odd}:{k}{ostatni_odd})" if k in "CDE" else (
            f"=IFERROR(C{r}/D{r},0)" if k == "F" else f"=SUM(G{pierwszy_odd}:G{ostatni_odd})")
        ws[f"{k}{r}"].number_format = fmt
    styl_wiersza(ws, r, 2, 7, pogrub=True)

    # Miesiące x oddziały
    w = r + 3
    ws.cell(w, 2, "Przychód miesięczny wg oddziałów").font = f_sekcja
    miesiace = sorted(dane["data"].dt.strftime("%Y-%m").unique())
    naglowek(ws, w + 1, 2, ["Miesiąc"] + oddzialy + ["Razem"])
    pierwszy_mies = w + 2
    for j, m in enumerate(miesiace):
        rr = pierwszy_mies + j
        ws.cell(rr, 2, m)
        for i in range(len(oddzialy)):
            kol = get_column_letter(3 + i)
            ws[f"{kol}{rr}"] = f'=SUMIFS({WART},{MIES},$B{rr},{ODDZ},{kol}${w + 1})'
            ws[f"{kol}{rr}"].number_format = ZL0
        kol_raz = get_column_letter(3 + len(oddzialy))
        ws[f"{kol_raz}{rr}"] = f"=SUM(C{rr}:{get_column_letter(2 + len(oddzialy))}{rr})"
        ws[f"{kol_raz}{rr}"].number_format = ZL0
        styl_wiersza(ws, rr, 2, 3 + len(oddzialy))
        ws[f"{kol_raz}{rr}"].font = f_bold
    ostatni_mies = pierwszy_mies + len(miesiace) - 1

    wykres = LineChart()
    wykres.title = "Przychód miesięczny"
    wykres.y_axis.title = "zł"
    wykres.y_axis.numFmt = '#,##0'
    wykres.height, wykres.width = 7.5, 17
    wykres.add_data(Reference(ws, min_col=3, max_col=2 + len(oddzialy), min_row=w + 1, max_row=ostatni_mies),
                    titles_from_data=True)
    wykres.set_categories(Reference(ws, min_col=2, min_row=pierwszy_mies, max_row=ostatni_mies))
    wykres.legend.position = "b"
    for seria, kolor in zip(wykres.series, [GRANAT, "E07A3F", "2E9E8F", "8E6BBF", "C9A227"]):
        seria.smooth = False
        seria.graphicalProperties.line.solidFill = kolor
        seria.graphicalProperties.line.width = 28000
    ws.add_chart(wykres, f"B{ostatni_mies + 2}")

    # Produkty (kolejność wg przychodu, wartości jako formuły)
    w = ostatni_mies + 18
    ws.row_breaks.append(Break(id=w - 1))  # ranking produktów na nowej stronie przy wydruku
    ws.cell(w, 2, "Ranking produktów").font = f_sekcja
    naglowek(ws, w + 1, 2, ["Produkt", "Przychód", "Sztuki", "Udział"])
    kolejnosc = dane.assign(w=dane["ilosc"] * dane["cena"]).groupby("produkt")["w"].sum().sort_values(ascending=False)
    pierwszy_prod = w + 2
    for j, p in enumerate(kolejnosc.index):
        rr = pierwszy_prod + j
        ws.cell(rr, 2, p)
        ws.cell(rr, 3, f'=SUMIFS({WART},{PROD},B{rr})').number_format = ZL0
        ws.cell(rr, 4, f'=SUMIFS({ILOSC},{PROD},B{rr})').number_format = SZT
        ws.cell(rr, 5, f'=IFERROR(C{rr}/$B$6,0)').number_format = PROC
        styl_wiersza(ws, rr, 2, 5)
    ostatni_prod = pierwszy_prod + len(kolejnosc) - 1

    slupki = BarChart()
    slupki.type = "bar"
    slupki.title = "Przychód wg produktu"
    slupki.legend = None
    slupki.x_axis.scaling.orientation = "maxMin"  # największy na górze
    slupki.y_axis.numFmt = '#,##0'
    slupki.y_axis.crosses = "max"  # oś wartości na dole przy odwróconej kolejności
    slupki.height, slupki.width = 7.5, 17
    slupki.add_data(Reference(ws, min_col=3, min_row=w + 1, max_row=ostatni_prod), titles_from_data=True)
    slupki.set_categories(Reference(ws, min_col=2, min_row=pierwszy_prod, max_row=ostatni_prod))
    slupki.series[0].graphicalProperties.solidFill = GRANAT
    ws.add_chart(slupki, f"B{ostatni_prod + 2}")

    # ---------------- Arkusz: Handlowcy ----------------
    wh.sheet_view.showGridLines = False
    wh.column_dimensions["A"].width = 2
    for k, szer in zip("BCDEFG", [6, 24, 14, 16, 14, 16]):
        wh.column_dimensions[k].width = szer
    wh["B2"] = "Ranking handlowców"
    wh["B2"].font = f_tytul
    naglowek(wh, 4, 2, ["#", "Handlowiec", "Oddział", "Przychód", "Transakcje", "Średni paragon"])
    hand = (dane.assign(w=dane["ilosc"] * dane["cena"]).groupby(["handlowiec", "oddzial"])["w"].sum()
            .sort_values(ascending=False))
    for j, (h, o) in enumerate(hand.index):
        rr = 5 + j
        wh.cell(rr, 2, j + 1)
        wh.cell(rr, 3, h)
        wh.cell(rr, 4, o)
        wh.cell(rr, 5, f'=SUMIFS({WART},{HAND},C{rr},{ODDZ},D{rr})').number_format = ZL0
        wh.cell(rr, 6, f'=COUNTIFS({HAND},C{rr},{ODDZ},D{rr})').number_format = SZT
        wh.cell(rr, 7, f'=IFERROR(E{rr}/F{rr},0)').number_format = ZL
        styl_wiersza(wh, rr, 2, 7, pogrub=(j < 3))
    wh.freeze_panes = "A5"

    # ---------------- Arkusz: Log czyszczenia ----------------
    wl.sheet_view.showGridLines = False
    wl.column_dimensions["A"].width = 2
    wl.column_dimensions["B"].width = 28
    wl.column_dimensions["C"].width = 58
    wl.column_dimensions["D"].width = 12
    wl["B2"] = "Co zostało poprawione w danych"
    wl["B2"].font = f_tytul
    wl["B3"] = "Pełna lista zmian – żeby było jasne, skąd biorą się liczby w raporcie."
    wl["B3"].font = f_pod
    naglowek(wl, 5, 2, ["Plik źródłowy", "Operacja", "Wierszy"])
    for j, (plik, co, ile) in enumerate(log):
        rr = 6 + j
        wl.cell(rr, 2, plik)
        wl.cell(rr, 3, co)
        wl.cell(rr, 4, ile).number_format = SZT
        styl_wiersza(wl, rr, 2, 4, pogrub=co.startswith("Wierszy po"))

    for arkusz in (ws, wh, wl, wd):
        arkusz.sheet_properties.tabColor = GRANAT
        # wydruk / PDF: cała szerokość na jednej stronie
        arkusz.page_setup.orientation = "portrait"
        arkusz.page_setup.fitToWidth = 1
        arkusz.page_setup.fitToHeight = 0
        arkusz.sheet_properties.pageSetUpPr.fitToPage = True
    wd.print_title_rows = "1:1"
    wb.save(PLIK_WYJSCIE)


def main():
    pliki = sorted(p for p in FOLDER_WEJSCIE.glob("*") if p.suffix.lower() in {".xlsx", ".xls", ".csv"})
    if not pliki:
        sys.exit(f"Brak plików w {FOLDER_WEJSCIE}. Uruchom najpierw generuj_dane_przykladowe.py")
    print(f"Wczytuję {len(pliki)} plik(i):")
    czesci = []
    for p in pliki:
        df = oczysc(p)
        print(f"  ✓ {p.name:30s} {len(df):>5} wierszy  ({df['oddzial'].iloc[0]})")
        czesci.append(df)
    dane = pd.concat(czesci, ignore_index=True)
    zbuduj_raport(dane)
    przychod = f"{(dane['ilosc'] * dane['cena']).sum():,.2f}".replace(",", " ").replace(".", ",")
    print(f"\nGotowe → {PLIK_WYJSCIE.name}  ({len(dane)} transakcji, przychód {przychod} zł)")


if __name__ == "__main__":
    main()
