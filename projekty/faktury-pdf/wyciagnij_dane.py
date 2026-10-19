"""
Faktury PDF → zestawienie w Excelu
==================================

Wrzuć faktury PDF do folderu `faktury/` i uruchom:

    python wyciagnij_dane.py

Skrypt dla każdej faktury:
  1. odczytuje tekst z PDF-a (różne układy od różnych dostawców),
  2. wyciąga: dostawcę, NIP, numer, datę, kwoty netto / VAT / brutto,
  3. sprawdza poprawność:
       - czy netto + VAT = brutto,
       - czy NIP ma poprawną sumę kontrolną,
       - czy faktura nie jest duplikatem,
       - czy PDF w ogóle ma tekst (skan wymaga OCR),
  4. zapisuje `zestawienie_faktur.xlsx` z listą faktur, podsumowaniem
     wg dostawców i miesięcy oraz listą faktur do ręcznego sprawdzenia.
"""
from __future__ import annotations

import re
import sys
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

FOLDER = Path(__file__).parent / "faktury"
WYJSCIE = Path(__file__).parent / "zestawienie_faktur.xlsx"

# ---------------------------------------------------------------------------
# Wzorce – dopisz nowy wariant etykiety, gdy pojawi się nowy dostawca
# ---------------------------------------------------------------------------
WZ_NUMER = [r"Faktura VAT nr\s+(\S+)", r"FAKTURA NR\s+(\S+)", r"Numer faktury:\s*(\S+)", r"Faktura nr\s+(\S+)"]
WZ_DATA = [r"Data wystawienia:\s*(.+)", r"Wystawiono:\s*(.+)", r"\bData:\s*(.+)"]
KWOTA = r"(\d{1,3}(?:[  ]\d{3})*(?:[.,]\d{2})|\d+[.,]\d{2})"
WZ_NETTO = [rf"Razem netto:\s*{KWOTA}", rf"Wartość netto:\s*{KWOTA}", rf"SUMA\s+netto\s+{KWOTA}"]
WZ_VAT = [rf"VAT \d+%:\s*{KWOTA}", rf"Kwota VAT[^:]*:\s*{KWOTA}", rf"\bVAT\s+{KWOTA}"]
WZ_BRUTTO = [rf"Razem brutto:\s*{KWOTA}", rf"Do zapłaty:\s*{KWOTA}", rf"brutto\s+{KWOTA}"]
MIESIACE = {m: i for i, m in enumerate(["stycznia", "lutego", "marca", "kwietnia", "maja", "czerwca", "lipca",
                                         "sierpnia", "września", "października", "listopada", "grudnia"], 1)}


@dataclass
class Faktura:
    plik: str
    dostawca: str = ""
    nip: str = ""
    numer: str = ""
    data: date | None = None
    netto: float | None = None
    vat: float | None = None
    brutto: float | None = None
    uwagi: list[str] = field(default_factory=list)

    @property
    def status(self) -> str:
        return "OK" if not self.uwagi else "DO SPRAWDZENIA"


def szukaj(wzorce, tekst):
    for w in wzorce:
        m = re.search(w, tekst)
        if m:
            return m.group(1).strip()
    return None


def kwota(s):
    if s is None:
        return None
    s = s.replace(" ", "").replace(" ", "").replace(",", ".")
    return float(s)


def parsuj_date(s):
    if not s:
        return None
    s = s.strip()
    if m := re.match(r"(\d{4})-(\d{2})-(\d{2})", s):
        return date(int(m[1]), int(m[2]), int(m[3]))
    if m := re.match(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", s):
        return date(int(m[3]), int(m[2]), int(m[1]))
    if m := re.match(r"(\d{1,2})\s+([a-ząćęłńóśźż]+)\s+(\d{4})", s):
        if m[2] in MIESIACE:
            return date(int(m[3]), MIESIACE[m[2]], int(m[1]))
    return None


def nip_poprawny(nip: str) -> bool:
    if not re.fullmatch(r"\d{10}", nip or ""):
        return False
    wagi = [6, 5, 7, 2, 3, 4, 5, 6, 7]
    return sum(int(c) * w for c, w in zip(nip, wagi)) % 11 == int(nip[9])


def odczytaj(plik: Path) -> Faktura:
    f = Faktura(plik=plik.name)
    with pdfplumber.open(plik) as pdf:
        strona = pdf.pages[0]
        tekst = strona.extract_text() or ""
        if len(tekst.strip()) < 20:
            f.uwagi.append("Brak tekstu w PDF (skan) – wymaga OCR albo ręcznego przepisania")
            return f
        # dane sprzedawcy są w lewej kolumnie – wycinamy lewą połowę strony
        lewa = strona.crop((0, 0, strona.width / 2, strona.height)).extract_text() or ""

    linie = lewa.splitlines()
    for i, l in enumerate(linie):
        if l.strip().startswith("Sprzedawca") and i + 1 < len(linie):
            f.dostawca = linie[i + 1].strip()
            break
    if m := re.search(r"NIP:?\s*([\d\- ]{10,13})", lewa):
        f.nip = re.sub(r"\D", "", m[1])

    f.numer = szukaj(WZ_NUMER, tekst) or ""
    f.data = parsuj_date(szukaj(WZ_DATA, tekst))
    f.netto = kwota(szukaj(WZ_NETTO, tekst))
    f.vat = kwota(szukaj(WZ_VAT, tekst))
    f.brutto = kwota(szukaj(WZ_BRUTTO, tekst))

    # --- walidacja ---
    for nazwa, wartosc in [("dostawcy", f.dostawca), ("numeru", f.numer), ("daty", f.data),
                           ("kwoty netto", f.netto), ("kwoty VAT", f.vat), ("kwoty brutto", f.brutto)]:
        if not wartosc and wartosc != 0:
            f.uwagi.append(f"Nie znaleziono {nazwa}")
    if f.nip and not nip_poprawny(f.nip):
        f.uwagi.append(f"Błędny NIP {f.nip} (zła suma kontrolna)")
    if None not in (f.netto, f.vat, f.brutto) and abs(f.netto + f.vat - f.brutto) > 0.01:
        f.uwagi.append(f"Netto + VAT = {f.netto + f.vat:.2f}, a na fakturze brutto {f.brutto:.2f}")
    return f


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
FONT = "Arial"
GRANAT = "1F3A5F"
H = Font(name=FONT, bold=True, color="FFFFFF")
N = Font(name=FONT, size=10)
B = Font(name=FONT, size=10, bold=True)
WYP_H = PatternFill("solid", fgColor=GRANAT)
ZL = '#,##0.00 "zł"'


def naglowek(ws, wiersz, teksty, kol=1):
    for i, t in enumerate(teksty):
        c = ws.cell(wiersz, kol + i, t)
        c.font, c.fill = H, WYP_H
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def zapisz(faktury: list[Faktura]):
    wb = Workbook()
    ws = wb.active
    ws.title = "Faktury"
    kol = ["Plik", "Dostawca", "NIP dostawcy", "Nr faktury", "Data", "Miesiąc", "Netto", "VAT", "Brutto",
           "Różnica (brutto − netto − VAT)", "Status", "Uwagi"]
    naglowek(ws, 1, kol)
    faktury = sorted(faktury, key=lambda f: (f.data or date.max, f.dostawca))
    for r, f in enumerate(faktury, start=2):
        vals = [f.plik, f.dostawca, f.nip, f.numer, f.data, f.data.strftime("%Y-%m") if f.data else "",
                f.netto, f.vat, f.brutto]
        for k, v in enumerate(vals, 1):
            ws.cell(r, k, v).font = N
        ws.cell(r, 5).number_format = "yyyy-mm-dd"
        for k in (7, 8, 9):
            ws.cell(r, k).number_format = ZL
        ws.cell(r, 10, f'=IF(COUNT(G{r}:I{r})=3,ROUND(I{r}-G{r}-H{r},2),"")').number_format = ZL
        ws.cell(r, 10).font = N
        ws.cell(r, 11, f.status).font = B
        ws.cell(r, 12, "; ".join(f.uwagi)).font = N
    last = len(faktury) + 1
    tab = Table(displayName="Faktury", ref=f"A1:L{last}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(tab)
    czerw = PatternFill("solid", fgColor="FDE2E1")
    ws.conditional_formatting.add(f"A2:L{last}", FormulaRule(formula=[f'$K2="DO SPRAWDZENIA"'], fill=czerw))
    for k, w in enumerate([24, 34, 14, 18, 12, 10, 13, 12, 13, 16, 17, 60], 1):
        ws.column_dimensions[get_column_letter(k)].width = w
    ws.freeze_panes = "A2"

    # --- Podsumowanie ---
    ps = wb.create_sheet("Podsumowanie")
    ps.sheet_view.showGridLines = False
    ps["A1"] = "Zestawienie kosztów z faktur"; ps["A1"].font = Font(name=FONT, size=16, bold=True, color=GRANAT)
    ps["A2"] = "Liczone formułami z arkusza „Faktury”. Kwoty obejmują tylko faktury ze statusem OK – faktury do sprawdzenia są policzone osobno."
    ps["A2"].font = Font(name=FONT, size=9, italic=True, color="666666")
    R = lambda c: f"Faktury!${c}$2:${c}${last}"  # noqa: E731
    naglowek(ps, 4, ["Dostawca", "Faktury OK", "Netto", "VAT", "Brutto", "Do sprawdzenia"])
    dostawcy = sorted({f.dostawca for f in faktury if f.dostawca})
    for j, d in enumerate(dostawcy):
        r = 5 + j
        ps.cell(r, 1, d).font = N
        ps.cell(r, 2, f'=COUNTIFS({R("B")},A{r},{R("K")},"OK")')
        for k, c in zip((3, 4, 5), "GHI"):
            ps.cell(r, k, f'=SUMIFS({R(c)},{R("B")},A{r},{R("K")},"OK")').number_format = ZL
        ps.cell(r, 6, f'=COUNTIFS({R("B")},A{r},{R("K")},"DO SPRAWDZENIA")')
        for k in range(2, 7):
            ps.cell(r, k).font = N
    r = 5 + len(dostawcy)
    ps.cell(r, 1, "Razem").font = B
    for k in range(2, 7):
        cl = get_column_letter(k)
        ps.cell(r, k, f"=SUM({cl}5:{cl}{r - 1})").font = B
        if k in (3, 4, 5):
            ps.cell(r, k).number_format = ZL
    # razem do sprawdzenia liczymy ze wszystkich faktur – także skanów bez rozpoznanego dostawcy
    ps.cell(r, 6, f'=COUNTIFS({R("K")},"DO SPRAWDZENIA")').font = B

    r += 3
    ps.cell(r - 1, 1, "Koszty brutto wg miesięcy (faktury OK)").font = Font(name=FONT, size=12, bold=True, color=GRANAT)
    naglowek(ps, r, ["Miesiąc"] + dostawcy + ["Razem"])
    miesiace = sorted({f.data.strftime("%Y-%m") for f in faktury if f.data})
    for j, m in enumerate(miesiace):
        rr = r + 1 + j
        ps.cell(rr, 1, m).font = N
        for k, d in enumerate(dostawcy, 2):
            ps.cell(rr, k, f'=SUMIFS({R("I")},{R("F")},$A{rr},{R("B")},{get_column_letter(k)}${r},{R("K")},"OK")').number_format = ZL
            ps.cell(rr, k).font = N
        kr = len(dostawcy) + 2
        ps.cell(rr, kr, f"=SUM(B{rr}:{get_column_letter(kr - 1)}{rr})").number_format = ZL
        ps.cell(rr, kr).font = B
    for k, w in enumerate([34, 16, 16, 16, 16, 16], 1):
        ps.column_dimensions[get_column_letter(k)].width = w

    # --- Do sprawdzenia ---
    ds = wb.create_sheet("Do sprawdzenia")
    ds["A1"] = "Faktury wymagające uwagi"; ds["A1"].font = Font(name=FONT, size=16, bold=True, color=GRANAT)
    naglowek(ds, 3, ["Plik", "Dostawca", "Nr faktury", "Problem"])
    rr = 4
    for f in faktury:
        for u in f.uwagi:
            for k, v in enumerate([f.plik, f.dostawca, f.numer, u], 1):
                ds.cell(rr, k, v).font = N
            rr += 1
    if rr == 4:
        ds.cell(4, 1, "Wszystkie faktury są poprawne.").font = N
    for k, w in enumerate([26, 34, 18, 70], 1):
        ds.column_dimensions[get_column_letter(k)].width = w

    for arkusz in wb.worksheets:  # wydruk: cała szerokość na jednej stronie
        arkusz.page_setup.orientation = "landscape"
        arkusz.page_setup.fitToWidth, arkusz.page_setup.fitToHeight = 1, 0
        arkusz.sheet_properties.pageSetUpPr.fitToPage = True
    wb.move_sheet("Podsumowanie", offset=-1)
    wb.active = 0
    wb.save(WYJSCIE)


def main():
    pliki = sorted(FOLDER.glob("*.pdf"))
    if not pliki:
        sys.exit(f"Brak plików PDF w {FOLDER}. Uruchom najpierw generuj_faktury.py")
    faktury = [odczytaj(p) for p in pliki]

    # duplikaty: ten sam dostawca i numer w więcej niż jednym pliku
    widziane = {}
    for f in faktury:
        if not f.numer:
            continue
        klucz = (f.nip, f.numer)
        if klucz in widziane:
            f.uwagi.append(f"Duplikat – ta sama faktura co w pliku {widziane[klucz]}")
        else:
            widziane[klucz] = f.plik

    zapisz(faktury)
    ok = sum(f.status == "OK" for f in faktury)
    print(f"Przetworzono {len(faktury)} faktur: {ok} OK, {len(faktury) - ok} do sprawdzenia")
    for f in faktury:
        for u in f.uwagi:
            print(f"  ! {f.plik}: {u}")
    print(f"Gotowe → {WYJSCIE.name}")


if __name__ == "__main__":
    main()
